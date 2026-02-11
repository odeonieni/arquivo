from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, NamedStyle, numbers
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

cores = ["Branco", "Preto", "Grafite", "Laranja", "Azul", "Cinza", "Amarelo"]

wb = Workbook()

# Estilos
date_style = NamedStyle(name="date_style")
date_style.number_format = numbers.FORMAT_DATE_YYYYMMDD2
wb.add_named_style(date_style)

kg_style = NamedStyle(name="kg_style")
kg_style.number_format = "0.00"
wb.add_named_style(kg_style)

header_fill = PatternFill("solid", fgColor="D9EAF7")
header_font = Font(bold=True)

def make_header(ws, headers):
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(headers[col_idx - 1]) + 2)

# 1) Aba Cores
ws_cores = wb.active
ws_cores.title = "Cores"
make_header(ws_cores, ["Cor"])
for c in cores:
    ws_cores.append([c])
ws_cores["C1"] = "Unidade: kg"
ws_cores["C1"].font = Font(italic=True, color="777777")

# 2) Aba Entradas
ws_ent = wb.create_sheet("Entradas")
make_header(ws_ent, ["Data", "Cor", "Qtd (kg)"])
ws_ent.column_dimensions["A"].width = 14
ws_ent.column_dimensions["B"].width = 16
ws_ent.column_dimensions["C"].width = 12

# Validação de dados para Cor
dv_cor_ent = DataValidation(type="list", formula1="=Cores!$A$2:$A$8", allow_blank=True)
ws_ent.add_data_validation(dv_cor_ent)
dv_cor_ent.add(f"B2:B5000")

# Estilos de data e kg
for r in range(2, 2000):
    ws_ent.cell(row=r, column=1).style = "date_style"
    ws_ent.cell(row=r, column=3).style = "kg_style"

# 3) Aba Estoque_Mensal
ws_est = wb.create_sheet("Estoque_Mensal")
make_header(ws_est, ["Mês", "Cor", "Estoque Final (kg)", "Estoque Inicial (kg) (opcional)"])
ws_est.column_dimensions["A"].width = 14
ws_est.column_dimensions["B"].width = 16
ws_est.column_dimensions["C"].width = 18
ws_est.column_dimensions["D"].width = 28

# Validação de dados para Cor
dv_cor_est = DataValidation(type="list", formula1="=Cores!$A$2:$A$8", allow_blank=True)
ws_est.add_data_validation(dv_cor_est)
dv_cor_est.add(f"B2:B5000")

# Estilos
for r in range(2, 2000):
    ws_est.cell(row=r, column=1).style = "date_style"
    ws_est.cell(row=r, column=3).style = "kg_style"
    ws_est.cell(row=r, column=4).style = "kg_style"

# 4) Aba Resumo
ws_res = wb.create_sheet("Resumo")
make_header(ws_res, ["Mês", "Cor", "Inicial (kg)", "Entradas (kg)", "Final (kg)", "Consumo (kg)"])
ws_res.column_dimensions["A"].width = 14
ws_res.column_dimensions["B"].width = 16
ws_res.column_dimensions["C"].width = 14
ws_res.column_dimensions["D"].width = 16
ws_res.column_dimensions["E"].width = 12
ws_res.column_dimensions["F"].width = 14

# Estilos das colunas numéricas
for r in range(2, 2000):
    ws_res.cell(row=r, column=1).style = "date_style"
    ws_res.cell(row=r, column=3).style = "kg_style"
    ws_res.cell(row=r, column=4).style = "kg_style"
    ws_res.cell(row=r, column=5).style = "kg_style"
    ws_res.cell(row=r, column=6).style = "kg_style"

# Fórmulas (armazenadas em inglês)
for r in range(2, 502):
    ws_res.cell(row=r, column=3).value = (
        '=IFERROR('
        'INDEX(Estoque_Mensal!$C:$C,'
        'SUMPRODUCT((Estoque_Mensal!$B:$B=B{row})*(Estoque_Mensal!$A:$A=EOMONTH(A{row},-1))*ROW(Estoque_Mensal!$B:$B))'
        '),'
        'IFERROR('
        'INDEX(Estoque_Mensal!$D:$D,'
        'SUMPRODUCT((Estoque_Mensal!$B:$B=B{row})*(Estoque_Mensal!$A:$A=A{row})*ROW(Estoque_Mensal!$B:$B))'
        '),""'
        ')'
        ')'
    ).format(row=r)

    ws_res.cell(row=r, column=4).value = (
        '=IFERROR('
        'SUMIFS(Entradas!$C:$C,Entradas!$B:$B,B{row},Entradas!$A:$A,">="&DATE(YEAR(A{row}),MONTH(A{row}),1),Entradas!$A:$A,"<="&EOMONTH(A{row},0))'
        ',0)'
    ).format(row=r)

    ws_res.cell(row=r, column=5).value = (
        '=IFERROR('
        'INDEX(Estoque_Mensal!$C:$C,'
        'SUMPRODUCT((Estoque_Mensal!$B:$B=B{row})*(Estoque_Mensal!$A:$A=A{row})*ROW(Estoque_Mensal!$B:$B))'
        '),""'
        ')'
    ).format(row=r)

    ws_res.cell(row=r, column=6).value = '=C{row}+D{row}-E{row}'.format(row=r)

wb.save("Controle_Consumo_Tinta.xlsx")
print("Arquivo gerado: Controle_Consumo_Tinta.xlsx")
